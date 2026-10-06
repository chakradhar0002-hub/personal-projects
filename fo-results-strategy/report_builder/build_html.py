"""Write a single-file HTML viewer of the results workbook (values as recalculated), plus an optional cross-check tab.

    python3 build_html.py ../reports/results_report_last_15_quarters.xlsx ../reports/results_report_last_15_quarters.html \
        [cross_check_summary.json]
"""
import json, sys
from datetime import date, datetime
from openpyxl import load_workbook

SRC, OUT = sys.argv[1], sys.argv[2]
CHECK = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else None
wb = load_workbook(SRC, data_only=True, read_only=True)


def fmt_code(number_format):
    """Compact format code for the page: p1 / p2 = percent with 1 or 2 decimals, n0 / n1 / n2 = number, d = date, t = text."""
    f = number_format or "General"
    if "%" in f:
        return "p2" if "0.00%" in f else "p1"
    if "yy" in f or "mmm" in f:
        return "d"
    if f == "@":
        return "t"
    if "0.00" in f:
        return "n2"
    if "0.0" in f:
        return "n1"
    if "0" in f:
        return "n0"
    return "t"


def cell_value(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 6)
    return v


def read_table(ws, header_row, first_row, section_row=None):
    rows = list(ws.iter_rows(min_row=1))
    header = [c.value for c in rows[header_row - 1]]
    while header and header[-1] is None:
        header.pop()
    n = len(header)
    sections = []
    if section_row:
        cur = ""
        for c in rows[section_row - 1][:n]:
            cur = c.value or cur
            sections.append(cur)
    fmts, data = ["t"] * n, []
    for r in rows[first_row - 1:]:
        vals = [cell_value(c.value) for c in r[:n]]
        if all(v in (None, "") for v in vals):
            continue
        for j, c in enumerate(r[:n]):
            if c.value is not None and fmts[j] == "t":
                code = fmt_code(c.number_format)
                fmts[j] = "t" if isinstance(c.value, str) else code
        data.append(vals)
    return {"title": rows[0][0].value, "note": rows[1][0].value, "cols": header, "fmts": fmts, "sections": sections,
            "rows": data}


tables = {
    "summary": read_table(wb["Summary"], 3, 4),
    "industry": read_table(wb["By industry"], 3, 4),
    "stock": read_table(wb["By stock"], 3, 4),
    "events": read_table(wb["Events"], 5, 6, section_row=4),
}
notes = [[r[0], r[1]] for r in wb["Notes"].iter_rows(min_row=3, values_only=True) if r[0]]
payload = json.dumps({"tables": tables, "notes": notes, "check": CHECK}, separators=(",", ":"), ensure_ascii=False)
payload = payload.replace("</", "<\\/")

HTML = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>F&amp;O Results Report</title>
<style>
:root{--bg:#f7f7f5;--panel:#fff;--ink:#1d1d1b;--muted:#6b6b66;--line:#e3e2dc;--head:#f0efe9;--accent:#2f5d8a;
--pos:#1e7b3c;--neg:#b3261e;--posbg:#e6f4ea;--negbg:#fce8e6;--chip:#eceae3;--chipon:#2f5d8a;--chipink:#fff}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#161615;--panel:#1f1f1d;--ink:#ecebe6;--muted:#a3a29c;
--line:#34332f;--head:#272724;--accent:#8db7e0;--pos:#7fd39a;--neg:#f2a19a;--posbg:#1d3325;--negbg:#3a2220;--chip:#2c2b28;
--chipon:#8db7e0;--chipink:#111}}
:root[data-theme="dark"]{--bg:#161615;--panel:#1f1f1d;--ink:#ecebe6;--muted:#a3a29c;--line:#34332f;--head:#272724;--accent:#8db7e0;
--pos:#7fd39a;--neg:#f2a19a;--posbg:#1d3325;--negbg:#3a2220;--chip:#2c2b28;--chipon:#8db7e0;--chipink:#111}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:13px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
header{padding:16px 16px 0;max-width:100%}
h1{font-size:18px;margin:0 0 4px;padding-right:72px}
.sub{color:var(--muted);margin:0 0 12px;max-width:110ch}
nav{display:flex;gap:4px;flex-wrap:wrap;border-bottom:1px solid var(--line);padding:0 16px}
nav button{border:0;background:none;color:var(--muted);padding:8px 12px;font:inherit;cursor:pointer;border-bottom:2px solid transparent}
nav button.on{color:var(--ink);border-bottom-color:var(--accent);font-weight:600}
main{padding:12px 16px 32px}
.bar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-bottom:10px}
.bar input,.bar select{font:inherit;color:var(--ink);background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:5px 8px}
.bar input{min-width:180px}
.chips{display:flex;gap:6px;flex-wrap:wrap;margin-bottom:10px}
.chip{border:1px solid var(--line);background:var(--chip);color:var(--ink);border-radius:999px;padding:3px 10px;font:inherit;font-size:12px;cursor:pointer}
.chip.on{background:var(--chipon);color:var(--chipink);border-color:var(--chipon)}
.count{color:var(--muted);margin-left:auto}
.wrap{overflow:auto;max-height:calc(100vh - 210px);border:1px solid var(--line);border-radius:8px;background:var(--panel)}
table{border-collapse:separate;border-spacing:0;font-variant-numeric:tabular-nums;width:max-content;min-width:100%}
th,td{padding:5px 8px;border-bottom:1px solid var(--line);white-space:nowrap;text-align:right}
th{position:sticky;top:0;background:var(--head);z-index:2;font-weight:600;cursor:pointer;vertical-align:bottom;white-space:normal;
max-width:110px;min-width:64px;line-height:1.25}
tr.sec th{top:0;z-index:3;cursor:default;text-align:left;color:var(--muted);font-weight:500;font-size:11px;max-width:none}
tr.sec+tr th{top:25px}
th.sorted{color:var(--accent)}
td.t,th.t{text-align:left}
td.k,th.k{position:sticky;left:0;background:var(--panel);z-index:1}
th.k{z-index:4;background:var(--head)}
td.pos{color:var(--pos)}td.neg{color:var(--neg)}
td.big.pos{background:var(--posbg)}td.big.neg{background:var(--negbg)}
tr:hover td{background:color-mix(in srgb,var(--accent) 7%,var(--panel))}
tr.total td{font-weight:600;border-top:2px solid var(--line)}
.pager{display:flex;gap:8px;align-items:center;margin-top:8px}
.pager button,.bar button{font:inherit;border:1px solid var(--line);background:var(--panel);color:var(--ink);border-radius:6px;padding:4px 10px;cursor:pointer}
.notes dt{font-weight:600;margin-top:10px}.notes dd{margin:2px 0 0;color:var(--muted);max-width:110ch}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:10px;margin-bottom:14px}
.card{background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:10px 12px}
.card b{display:block;font-size:20px}.card span{color:var(--muted)}
h2{font-size:15px;margin:18px 0 8px}
p.lead{max-width:110ch;margin:0 0 10px}
.theme{position:absolute;top:14px;right:16px}
@media (max-width:640px){.wrap{max-height:none}.bar input{min-width:0;flex:1}}
</style>
</head>
<body>
<header>
<button class="chip theme" id="theme" title="Switch light / dark">Theme</button>
<h1>F&amp;O stocks: results report, last 15 quarters</h1>
<p class="sub" id="sub"></p>
</header>
<nav id="tabs"></nav>
<main id="main"></main>
<script id="data" type="application/json">__DATA__</script>
<script>
const D = JSON.parse(document.getElementById("data").textContent);
const T = D.tables;
document.getElementById("sub").textContent = T.events.note || "";
const $ = (tag, attrs = {}, ...kids) => { const e = document.createElement(tag);
  for (const [k, v] of Object.entries(attrs)) { if (k === "class") e.className = v; else if (k.startsWith("on")) e.addEventListener(k.slice(2), v); else e.setAttribute(k, v); }
  for (const k of kids) e.append(k); return e; };
const MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"];
function fmt(v, f) {
  if (v === null || v === undefined || v === "") return "";
  if (typeof v !== "number") { if (f === "d" && /^\d{4}-\d{2}-\d{2}$/.test(v)) { const [y,m,d] = v.split("-"); return `${d}-${MON[+m-1]}-${y.slice(2)}`; } return String(v); }
  if (f === "p1") return (v * 100).toFixed(1) + "%";
  if (f === "p2") return (v * 100).toFixed(2) + "%";
  if (f === "n2") return v.toLocaleString("en-IN", {minimumFractionDigits: 2, maximumFractionDigits: 2});
  if (f === "n1") return v.toLocaleString("en-IN", {minimumFractionDigits: 1, maximumFractionDigits: 1});
  if (f === "n0") return v.toLocaleString("en-IN", {maximumFractionDigits: 0});
  return String(v);
}
// green / red only for changes and returns, not for shares, levels or ratios
const SIGNED = h => !/^%|^within|^same|median gap|comparisons|win rate|size of|^IV|IV before|IV after|IV \d|expected move|margin %$|^ROE|NPA|^ROA|range|expected/i.test(h);
function cls(v, f, key, signed, fill) {
  let c = (f === "t" || f === "d" || typeof v !== "number") ? "t" : "";
  if (signed && typeof v === "number" && (f === "p1" || f === "p2")) { c += v > 0 ? " pos" : v < 0 ? " neg" : ""; if (fill && Math.abs(v) >= 0.05) c += " big"; }
  if (key) c += " k"; return c;
}
// generic sortable table
function grid(tbl, opts = {}) {
  const st = {sort: null, dir: 1, page: 0, rows: tbl.rows}; const signed = tbl.cols.map(h => SIGNED(String(h)));
  const box = $("div"); const wrap = $("div", {class: "wrap"}); const pager = $("div", {class: "pager"});
  box.append(wrap, pager);
  const vis = () => (opts.visible ? opts.visible() : tbl.cols.map((_, i) => i));
  function draw() {
    let rows = opts.filter ? tbl.rows.filter(opts.filter) : tbl.rows.slice();
    const pinned = opts.pinLast ? rows.filter(opts.pinLast) : []; if (opts.pinLast) rows = rows.filter(r => !opts.pinLast(r));
    if (st.sort !== null) { const j = st.sort; rows.sort((a, b) => { const x = a[j], y = b[j];
      if (x === null || x === "" || x === undefined) return 1; if (y === null || y === "" || y === undefined) return -1;
      return (x > y ? 1 : x < y ? -1 : 0) * st.dir; }); }
    st.count = rows.length; if (opts.onCount) opts.onCount(rows);
    const per = opts.perPage || 1e9, pages = Math.max(1, Math.ceil(rows.length / per)); st.page = Math.min(st.page, pages - 1);
    const shown = rows.slice(st.page * per, (st.page + 1) * per).concat(pinned);
    const cols = vis(); const keyCol = opts.keyCol ?? 0;
    const t = $("table"); const thead = $("thead");
    if (tbl.sections && tbl.sections.length) { const tr = $("tr", {class: "sec"}); let last = null, span = null;
      for (const j of cols) { const s = tbl.sections[j]; if (s !== last) { span = $("th", {class: j === keyCol ? "t k" : "t", colspan: 1}, s || ""); tr.append(span); last = s; } else span.colSpan++; }
      thead.append(tr); }
    const hr = $("tr");
    for (const j of cols) { const f = tbl.fmts[j];
      const th = $("th", {class: (f === "t" || f === "d" ? "t" : "") + (j === keyCol ? " k" : "") + (st.sort === j ? " sorted" : ""), title: "Sort"},
        tbl.cols[j] + (st.sort === j ? (st.dir > 0 ? " ▲" : " ▼") : ""));
      th.onclick = () => { if (st.sort === j) st.dir = -st.dir; else { st.sort = j; st.dir = (f === "t" || f === "d") ? 1 : -1; } draw(); };
      hr.append(th); }
    thead.append(hr); t.append(thead);
    const tb = $("tbody");
    for (const r of shown) { const tr = $("tr", opts.pinLast && opts.pinLast(r) ? {class: "total"} : {});
      for (const j of cols) tr.append($("td", {class: cls(r[j], tbl.fmts[j], j === keyCol, signed[j], opts.fill)}, fmt(r[j], tbl.fmts[j])));
      tb.append(tr); }
    t.append(tb); wrap.replaceChildren(t);
    pager.replaceChildren();
    if (pages > 1) { pager.append($("button", {onclick: () => { st.page = Math.max(0, st.page - 1); draw(); }}, "‹ Prev"),
      $("span", {}, `Page ${st.page + 1} of ${pages}`), $("button", {onclick: () => { st.page = Math.min(pages - 1, st.page + 1); draw(); }}, "Next ›")); }
  }
  box.redraw = (resetPage) => { if (resetPage) st.page = 0; draw(); }; draw(); return box;
}
function simpleView(key, pinWord) {
  const tbl = T[key]; const v = $("div");
  v.append($("p", {class: "lead"}, tbl.note || ""));
  const search = $("input", {placeholder: "Filter rows…", type: "search"});
  const count = $("span", {class: "count"});
  v.append($("div", {class: "bar"}, search, count));
  const g = grid(tbl, {pinLast: r => pinWord && String(r[0]).startsWith(pinWord),
    filter: r => !search.value || r.slice(0, 3).join(" ").toLowerCase().includes(search.value.toLowerCase()),
    onCount: rows => count.textContent = `${rows.length} rows`});
  search.oninput = () => g.redraw(true); v.append(g); return v;
}
function eventsView() {
  const tbl = T.events; const v = $("div");
  const ci = name => tbl.cols.indexOf(name);
  const iQ = ci("Quarter"), iInd = ci("Industry"), iTim = ci("Timing"), iRule = ci("Run-up rule pass");
  const uniq = j => [...new Set(tbl.rows.map(r => r[j]).filter(x => x !== null && x !== ""))];
  const sel = (label, vals) => { const s = $("select", {}, $("option", {value: ""}, label)); for (const x of vals) s.append($("option", {value: x}, x)); return s; };
  const q = sel("All quarters", uniq(iQ)), ind = sel("All industries", uniq(iInd).sort()), tim = sel("All timings", uniq(iTim).sort());
  const rule = sel("Run-up rule: any", ["Yes", "No"]);
  const search = $("input", {placeholder: "Symbol or company…", type: "search"});
  const count = $("span", {class: "count"});
  const csvBtn = $("button", {}, "Download CSV (filtered)");
  v.append($("div", {class: "bar"}, search, q, ind, tim, rule, csvBtn, count));
  const secs = [...new Set(tbl.sections)]; const on = new Set(["Stock & quarter", "Results timing", "3-day window (Day -1, Result day, Day +1)", "Reaction day"]);
  const chips = $("div", {class: "chips"});
  for (const s of secs) { const c = $("button", {class: "chip" + (on.has(s) ? " on" : "")}, s.replace(/ \(.*\)$/, ""));
    c.title = s; c.onclick = () => { on.has(s) ? on.delete(s) : on.add(s); c.classList.toggle("on"); g.redraw(); }; chips.append(c); }
  const all = $("button", {class: "chip"}, "Show all columns"); all.onclick = () => { secs.forEach(s => on.add(s)); chips.querySelectorAll(".chip").forEach(c => c.classList.add("on")); g.redraw(); };
  chips.append(all); v.append(chips);
  const always = new Set([ci("Quarter"), ci("Symbol")]);
  const filter = r => (!q.value || r[iQ] === q.value) && (!ind.value || r[iInd] === ind.value) && (!tim.value || r[iTim] === tim.value) &&
    (!rule.value || r[iRule] === rule.value) &&
    (!search.value || (String(r[ci("Symbol")]) + " " + String(r[ci("Company")])).toLowerCase().includes(search.value.toLowerCase()));
  let current = [];
  const g = grid(tbl, {perPage: 100, keyCol: ci("Symbol"), filter, fill: true,
    visible: () => tbl.cols.map((_, j) => j).filter(j => always.has(j) || on.has(tbl.sections[j])),
    onCount: rows => { current = rows; count.textContent = `${rows.length.toLocaleString("en-IN")} of ${tbl.rows.length.toLocaleString("en-IN")} results`; }});
  for (const el of [q, ind, tim, rule]) el.onchange = () => g.redraw(true);
  search.oninput = () => g.redraw(true);
  csvBtn.onclick = () => { const esc = x => { const s = x === null || x === undefined ? "" : String(x); return /[",\n]/.test(s) ? '"' + s.replace(/"/g, '""') + '"' : s; };
    const lines = [tbl.cols.map(esc).join(",")].concat(current.map(r => r.map(esc).join(",")));
    const a = $("a", {href: URL.createObjectURL(new Blob([lines.join("\n")], {type: "text/csv"})), download: "results_events_filtered.csv"}); a.click(); };
  v.append(g); return v;
}
function notesView() {
  const dl = $("dl", {class: "notes"}); for (const [k, t] of D.notes) dl.append($("dt", {}, k), $("dd", {}, t)); return dl;
}
function checkView() {
  const c = D.check; const v = $("div");
  v.append($("p", {class: "lead"}, c.intro));
  const cards = $("div", {class: "cards"}); for (const [k, t] of c.cards) cards.append($("div", {class: "card"}, $("b", {}, k), $("span", {}, t)));
  v.append(cards);
  for (const s of c.sections) { v.append($("h2", {}, s.title)); if (s.text) v.append($("p", {class: "lead"}, s.text));
    if (s.table) v.append(grid(s.table, {})); }
  return v;
}
const VIEWS = [["Summary by quarter", () => simpleView("summary", "All")], ["By industry", () => simpleView("industry", "All")],
  ["By stock", () => simpleView("stock")], ["All results (rows)", eventsView]];
if (D.check) VIEWS.push(["BSE cross-check", checkView]);
VIEWS.push(["Notes", notesView]);
const nav = document.getElementById("tabs"), main = document.getElementById("main"); const cache = {};
function show(i) { nav.querySelectorAll("button").forEach((b, j) => b.classList.toggle("on", i === j));
  cache[i] = cache[i] || VIEWS[i][1](); main.replaceChildren(cache[i]); try { localStorage.setItem("tab", i); } catch (e) {} }
VIEWS.forEach(([name], i) => nav.append($("button", {onclick: () => show(i)}, name)));
let start = 0; try { start = +(localStorage.getItem("tab") || 0); } catch (e) {} show(Math.min(start, VIEWS.length - 1));
document.getElementById("theme").onclick = () => { const r = document.documentElement;
  const dark = r.dataset.theme ? r.dataset.theme === "dark" : matchMedia("(prefers-color-scheme: dark)").matches; r.dataset.theme = dark ? "light" : "dark"; };
</script>
</body>
</html>
"""
open(OUT, "w").write(HTML.replace("__DATA__", payload))
print("wrote", OUT, f"{len(HTML) + len(payload):,} bytes", {k: len(v["rows"]) for k, v in tables.items()})
